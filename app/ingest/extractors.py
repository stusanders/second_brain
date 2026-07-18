"""Source-type extractors: turn a dropped file / URL / pasted text into plain
text plus metadata. Deterministic code only — no model calls here (cost
discipline: extraction is not a judgment call).

Day-one types per build spec rev 2. Email ingest is deferred.
"""

import io
from dataclasses import dataclass

import httpx


@dataclass
class ExtractedSource:
    source_type: str
    source_ref: str  # filename or URL (display label; not a blob path)
    text: str
    raw_bytes: bytes = b""  # verbatim original, stored immutably on ingest
    filename: str = ""  # used to name the blob when raw_bytes is stored


class ExtractionError(Exception):
    pass


MAX_CHARS = 200_000


def extract_file(filename: str, data: bytes) -> ExtractedSource:
    name = filename.lower()
    if name.endswith(".pdf"):
        return ExtractedSource("pdf", filename, _pdf_text(data), data, filename)
    if name.endswith(".docx"):
        return ExtractedSource("docx", filename, _docx_text(data), data, filename)
    if name.endswith(".pptx"):
        return ExtractedSource("pptx", filename, _pptx_text(data), data, filename)
    if name.endswith((".xlsx", ".xlsm")):
        return ExtractedSource("xlsx", filename, _xlsx_text(data), data, filename)
    if name.endswith(".csv"):
        return ExtractedSource("csv", filename, _decode(data), data, filename)
    if name.endswith((".txt", ".md", ".markdown")):
        return ExtractedSource("text", filename, _decode(data), data, filename)
    raise ExtractionError(f"Unsupported file type: {filename}")


def extract_url(url: str) -> ExtractedSource:
    resp = httpx.get(url, follow_redirects=True, timeout=30, headers={"User-Agent": "llm-wiki-poc"})
    resp.raise_for_status()
    content_type = resp.headers.get("content-type", "")
    filename = url.rsplit("/", 1)[-1] or "source"
    if "pdf" in content_type or url.lower().endswith(".pdf"):
        return ExtractedSource("pdf_url", url, _pdf_text(resp.content), resp.content, filename)
    import trafilatura

    text = trafilatura.extract(resp.text, url=url)
    if not text:
        raise ExtractionError(f"Could not extract main content from {url}")
    return ExtractedSource("web", url, text[:MAX_CHARS], resp.content, f"{filename}.html")


def extract_pasted(text: str) -> ExtractedSource:
    body = text[:MAX_CHARS]
    return ExtractedSource("pasted_text", "pasted text", body, body.encode("utf-8"), "pasted.txt")


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")[:MAX_CHARS]


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        # Scanned PDF with no text layer. OCR is in scope per spec but needs
        # a system tesseract install; surface a clear error until wired up.
        raise ExtractionError("PDF has no extractable text (scanned?). OCR not yet wired up.")
    return text[:MAX_CHARS]


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)[:MAX_CHARS]


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Speaker notes] {notes}")
    return "\n".join(parts)[:MAX_CHARS]


def _xlsx_text(data: bytes) -> str:
    """Structure + a bounded sample, not a raw cell dump (per spec, the model
    summarizes structure and key figures)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"--- Sheet: {ws.title} ({ws.max_row} rows x {ws.max_column} cols) ---")
        for row in ws.iter_rows(max_row=min(ws.max_row or 0, 50), values_only=True):
            parts.append(" | ".join("" if v is None else str(v) for v in row))
        if (ws.max_row or 0) > 50:
            parts.append(f"... ({ws.max_row - 50} more rows not shown)")
    return "\n".join(parts)[:MAX_CHARS]
